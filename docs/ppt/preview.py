"""Rough Pillow renderer for the built deck — layout sanity check only.
Approximates fills, text wrapping and placement; not pixel-faithful.
Usage: python3 docs/ppt/preview.py -> docs/ppt/preview/slide-N.png
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

PX = 120  # px per inch
W, H = int(13.333 * PX), int(7.5 * PX)
SUP = "/System/Library/Fonts/Supplemental/"
FONTS = {
    ("Arial", False, False): SUP + "Arial.ttf",
    ("Arial", True, False): SUP + "Arial Bold.ttf",
    ("Arial", False, True): SUP + "Arial Italic.ttf",
    ("Arial", True, True): SUP + "Arial Bold Italic.ttf",
    ("Times New Roman", False, False): SUP + "Times New Roman.ttf",
    ("Times New Roman", True, False): SUP + "Times New Roman Bold.ttf",
    ("Times New Roman", False, True): SUP + "Times New Roman Italic.ttf",
    ("Times New Roman", True, True): SUP + "Times New Roman Bold Italic.ttf",
}
_font_cache = {}


def font(name, size_pt, bold, italic):
    px = max(8, int(size_pt * PX / 72))
    key = (name or "Arial", bool(bold), bool(italic), px)
    if key not in _font_cache:
        path = FONTS.get((key[0], key[1], key[2])) or FONTS[("Arial", *key[1:3])] \
            if (key[0], key[1], key[2]) not in FONTS else FONTS[key[:3]]
        try:
            _font_cache[key] = ImageFont.truetype(path, px)
        except Exception:
            _font_cache[key] = ImageFont.truetype(SUP + "Arial.ttf", px)
    return _font_cache[key]


def rgb_of(color_obj, default=(0, 0, 0)):
    try:
        c = color_obj.rgb
        return (c[0], c[1], c[2]) if c is not None else default
    except Exception:
        return default


def fill_rgb(shape):
    try:
        if shape.fill.type is not None and shape.fill.type == 1:
            return rgb_of(shape.fill.fore_color, None)
    except Exception:
        pass
    return None


def wrap(draw, text, fnt, width_px):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= width_px or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def draw_text(draw, shape, box):
    x0, y0, x1, y1 = box
    tf = shape.text_frame
    try:
        ml = Emu(tf.margin_left).inches * PX if tf.margin_left is not None else 0.1 * PX
        mr = Emu(tf.margin_right).inches * PX if tf.margin_right is not None else 0.1 * PX
        mt = Emu(tf.margin_top).inches * PX if tf.margin_top is not None else 0.05 * PX
        mb = Emu(tf.margin_bottom).inches * PX if tf.margin_bottom is not None else 0.05 * PX
    except Exception:
        ml = mr = 0.1 * PX
        mt = mb = 0.05 * PX
    y = y0 + mt
    for p in tf.paragraphs:
        runs = [(r.text, r.font) for r in p.runs if r.text]
        if not runs:
            y += 6
            continue
        sz = max([(rf.size.pt if rf.size else 11) for _, rf in runs])
        lh = sz * PX / 72 * 1.3
        full = "".join(t for t, _ in runs)
        rep = runs[0][1]
        fnt0 = font(rep.name, rep.size.pt if rep.size else 11,
                    rep.bold, rep.italic)
        col = rgb_of(rep.color, (20, 20, 20)) if rep.color and rep.color.type is not None else (20, 20, 20)
        lines = wrap(draw, full, fnt0, x1 - x0 - ml - mr)
        for ln in lines:
            lw = draw.textlength(ln, font=fnt0)
            if p.alignment is not None and int(p.alignment) in (2, 3):
                tx = x0 + ml + (x1 - x0 - ml - mr - lw) / 2
            elif p.alignment is not None and int(p.alignment) == 4:
                tx = x0 + ml
            else:
                tx = x0 + ml
            if y + lh > y1 - mb + 8:
                draw.text((tx, y), ln + " …", font=fnt0, fill=(200, 0, 0))
                return
            draw.text((tx, y), ln, font=fnt0, fill=col)
            y += lh
        sa = p.space_after.pt * PX / 72 if p.space_after else 0
        y += sa


def render_shape(img, draw, shape):
    st = shape.shape_type
    x0 = Emu(shape.left).inches * PX if shape.left is not None else 0
    y0 = Emu(shape.top).inches * PX if shape.top is not None else 0
    x1 = x0 + (Emu(shape.width).inches * PX if shape.width else 0)
    y1 = y0 + (Emu(shape.height).inches * PX if shape.height else 0)

    if st == MSO_SHAPE_TYPE.PICTURE:
        try:
            pim = Image.open(__import__("io").BytesIO(shape.image.blob))
            pim = pim.resize((int(x1 - x0), int(y1 - y0)))
            img.paste(pim, (int(x0), int(y0)))
        except Exception:
            draw.rectangle([x0, y0, x1, y1], outline=(180, 180, 180))
        return
    if shape.has_table if hasattr(shape, "has_table") else False:
        tbl = shape.table
        nrows = len(tbl.rows)
        rh = (y1 - y0) / nrows
        cws = [Emu(c.width).inches * PX for c in tbl.columns]
        cy = y0
        for ri, row in enumerate(tbl.rows):
            cx = x0
            for ci, cell in enumerate(row.cells):
                cw = cws[ci]
                try:
                    fc = cell.fill.fore_color.rgb
                    fill = (fc[0], fc[1], fc[2])
                except Exception:
                    fill = (255, 255, 255)
                draw.rectangle([cx, cy, cx + cw, cy + rh], fill=fill,
                               outline=(150, 150, 150))
                # cell text
                tf = cell.text_frame
                yy = cy + rh / 2 - 8
                for p in tf.paragraphs:
                    for r in p.runs:
                        fnt = font(r.font.name, r.font.size.pt if r.font.size else 10,
                                   r.font.bold, r.font.italic)
                        col = rgb_of(r.font.color, (20, 20, 20)) if r.font.color and r.font.color.type is not None else (20, 20, 20)
                        lines = wrap(draw, r.text, fnt, cw - 12)
                        for ln in lines:
                            draw.text((cx + 6, yy), ln, font=fnt, fill=col)
                            yy += (r.font.size.pt if r.font.size else 10) * PX / 72 * 1.25
                        break
                cx += cw
            cy += rh
        return

    f = fill_rgb(shape)
    name = getattr(shape, "name", "")
    if f:
        if "Oval" in name:
            draw.ellipse([x0, y0, x1, y1], fill=f, outline=(150, 150, 150))
        elif "CHEVRON" in str(getattr(shape, "auto_shape_type", "")):
            k = (x1 - x0) * 0.25
            pts = [(x0, y0), (x1 - k, y0), (x1, (y0 + y1) / 2),
                   (x1 - k, y1), (x0, y1), (x0 + k, (y0 + y1) / 2)]
            draw.polygon(pts, fill=f)
        elif "ROUND" in str(getattr(shape, "auto_shape_type", "")):
            draw.rounded_rectangle([x0, y0, x1, y1],
                                   radius=min(18, (y1 - y0) / 3),
                                   fill=f, outline=(170, 170, 170))
        else:
            draw.rectangle([x0, y0, x1, y1], fill=f)
    elif "Oval" in name:
        draw.ellipse([x0, y0, x1, y1], outline=(150, 150, 150), width=2)
    if shape.has_text_frame and shape.text_frame.text.strip():
        draw_text(draw, shape, (x0, y0, x1, y1))


prs = Presentation("/Users/abhinav/Projects/mantle/docs/ppt/Mantle_SIH26120.pptx")
os.makedirs("/Users/abhinav/Projects/mantle/docs/ppt/preview", exist_ok=True)
for i, s in enumerate(prs.slides):
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    for sh in s.shapes:
        try:
            render_shape(img, d, sh)
        except Exception as e:
            print(f"slide{i} shape {getattr(sh,'name','?')}: {e}")
    out = f"/Users/abhinav/Projects/mantle/docs/ppt/preview/slide-{i}.png"
    img.save(out)
    print("wrote", out)
